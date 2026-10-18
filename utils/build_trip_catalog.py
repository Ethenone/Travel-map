"""Build the static trip catalog from the workbooks in ../database.

Run from anywhere with: python utils/build_trip_catalog.py
Missing workbook cells remain null in the published catalog.
"""

from __future__ import annotations

import csv
import json
from collections import Counter
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from timezonefinder import TimezoneFinder


ROOT = Path(__file__).resolve().parents[1]
DATABASE = ROOT.parent / "database"
OUTPUT = ROOT / "data" / "trip_catalog.json"
TIMEZONES = TimezoneFinder(in_memory=True)


def value(item):
    if item is None or item == "":
        return None
    if isinstance(item, (datetime, date)):
        return item.strftime("%Y-%m-%d")
    if isinstance(item, time):
        return item.strftime("%H:%M")
    if isinstance(item, str):
        return item.strip() or None
    return item


def rows(filename, sheet):
    book = load_workbook(DATABASE / filename, read_only=True, data_only=True)
    iterator = book[sheet].iter_rows(values_only=True)
    headers = list(next(iterator))
    result = []
    for line in iterator:
        row = {key: value(line[index]) for index, key in enumerate(headers) if key and index < len(line)}
        result.append(row)
    book.close()
    return result


def csv_lookup(filename, key):
    with (ROOT / "reference-data" / filename).open(encoding="utf-8-sig", newline="") as handle:
        return {row[key]: row for row in csv.DictReader(handle) if row.get(key)}


STATIONS = csv_lookup("stations_data.csv", "车站")
AIRPORTS = csv_lookup("airports_data.csv", "iata")


def coordinates(record, lon="经度", lat="纬度"):
    try:
        return [round(float(record[lon]), 6), round(float(record[lat]), 6)]
    except (TypeError, ValueError, KeyError):
        return None


def station(name, city=None, country=None, foreign=None):
    record = (foreign or {}).get(name) or STATIONS.get(name) or {}
    point = coordinates(record, "lon", "lat") if foreign else coordinates(record)
    zone = TIMEZONES.timezone_at(lng=point[0], lat=point[1]) if point else None
    return {
        "name": name,
        "country": country or ("中国" if name in STATIONS else None),
        "province": record.get("省"),
        "city": city or record.get("市"),
        "timezone": zone,
        "coordinates": point,
    }


def duration_minutes(row):
    number = row.get("时长.1")
    if isinstance(number, (float, int)):
        return round(number)
    clock = row.get("时长") or row.get("飞行时间")
    if isinstance(clock, str) and ":" in clock:
        parts = clock.split(":")
        try:
            return int(parts[0]) * 60 + int(parts[1])
        except ValueError:
            pass
    return None


def utc_offset(day, clock, zone, day_offset=0):
    if not day or not zone:
        return None
    try:
        local = datetime.fromisoformat(f"{day}T{clock or '12:00'}") + timedelta(days=day_offset)
        minutes = round(local.replace(tzinfo=ZoneInfo(zone)).utcoffset().total_seconds() / 60)
    except (ValueError, KeyError, AttributeError):
        return None
    sign = "+" if minutes >= 0 else "-"
    hours, remainder = divmod(abs(minutes), 60)
    return f"UTC{sign}{hours}{f':{remainder:02d}' if remainder else ''}"


def shifted_date(day, day_offset):
    if not day:
        return None
    return (date.fromisoformat(day) + timedelta(days=day_offset)).isoformat()


def inferred_day_offset(day, departure_time, arrival_time, duration, departure_zone, arrival_zone):
    if not all((day, departure_time, arrival_time, departure_zone, arrival_zone)) or duration is None:
        return 0
    departure = datetime.fromisoformat(f"{day}T{departure_time}").replace(tzinfo=ZoneInfo(departure_zone))
    candidates = []
    for offset in range(-1, 3):
        arrival_day = shifted_date(day, offset)
        arrival = datetime.fromisoformat(f"{arrival_day}T{arrival_time}").replace(tzinfo=ZoneInfo(arrival_zone))
        elapsed = (arrival.astimezone(timezone.utc) - departure.astimezone(timezone.utc)).total_seconds() / 60
        if elapsed >= 0:
            candidates.append((abs(elapsed - duration), abs(offset), offset))
    return min(candidates)[2] if candidates else 0


def airport(code, name, english_name, city, country, city_provinces, day, clock, day_offset=0):
    record = AIRPORTS.get(code) or {}
    point = coordinates(record, "lon", "lat")
    province = city_provinces.get(city) if country == "China" else None
    zone = TIMEZONES.timezone_at(lng=point[0], lat=point[1]) if point else None
    return {
        "name": name or record.get("name") or code,
        "nameEn": english_name or record.get("name"),
        "code": code,
        "country": country or record.get("country"),
        "province": province,
        "city": city,
        "timezone": zone,
        "utcOffset": utc_offset(day, clock, zone, day_offset),
        "coordinates": point,
    }


def main():
    province_by_city = {}
    for record in STATIONS.values():
        city, province = record.get("市"), record.get("省")
        if city and province:
            province_by_city[city] = province
            province_by_city[city.removesuffix("市")] = province

    rarity = {}
    # Keep the summary sheet as a fallback for models absent from the manual list.
    book = load_workbook(DATABASE / "火车乘坐记录.xlsx", read_only=True, data_only=True)
    for line in book["车型"].iter_rows(min_row=2, values_only=True):
        if line[0] and len(line) > 3 and isinstance(line[3], str):
            rarity[str(line[0]).strip()] = line[3].strip()
    # Column D of 车型列表 contains the manually assigned rarity.
    for line in book["车型列表"].iter_rows(min_row=2, values_only=True):
        if line[0] and len(line) > 3 and isinstance(line[3], str):
            rarity[str(line[0]).strip()] = line[3].strip()
    book.close()
    train_rarity_totals = Counter(rarity.values())

    plane_rarity = {record["type"]: record.get("rarity") for record in rows("飞机乘坐记录.xlsx", "type") if record.get("type")}
    foreign_stations = {}
    for record in rows("境外铁路乘坐记录.xlsx", "车站位置"):
        if record.get("车站"):
            foreign_stations[record["车站"]] = record

    catalog = []
    for index, row in enumerate(rows("火车乘坐记录.xlsx", "乘坐列表"), 1):
        if not row.get("车次") or not row.get("上车站") or not row.get("下车站"):
            continue
        model = row.get("车型")
        model_key = str(model).strip() if model is not None else None
        model_rarity = rarity.get(model_key)
        origin = station(row["上车站"], row.get("上车城市"))
        destination = station(row["下车站"], row.get("下车城市"))
        day_offset = row.get("跨天") if isinstance(row.get("跨天"), int) else inferred_day_offset(row.get("日期"), row.get("上车时间"), row.get("下车时间"), duration_minutes(row), origin["timezone"], destination["timezone"])
        origin["utcOffset"] = utc_offset(row.get("日期"), row.get("上车时间"), origin["timezone"])
        destination["utcOffset"] = utc_offset(row.get("日期"), row.get("下车时间"), destination["timezone"], day_offset)
        catalog.append({
            "id": f"cn-rail-{index}", "mode": "train", "date": row.get("日期"),
            "sourceNumber": row.get("序号"),
            "departureDate": row.get("日期"), "arrivalDate": shifted_date(row.get("日期"), day_offset),
            "dayOffset": day_offset,
            "service": str(row["车次"]), "operator": row.get("局属"),
            "origin": origin, "destination": destination,
            "startTerminal": row.get("始发"), "endTerminal": row.get("终到"),
            "departureTime": row.get("上车时间"), "arrivalTime": row.get("下车时间"),
            "seat": row.get("座位"), "cabin": row.get("席位"),
            "durationMinutes": duration_minutes(row), "distanceKm": row.get("里程"), "speedKmh": row.get("速度"),
            "model": model, "rarity": model_rarity,
            "rarityTotal": train_rarity_totals.get(model_rarity), "routeId": index,
        })

    for index, row in enumerate(rows("境外铁路乘坐记录.xlsx", "乘坐列表"), 1):
        if not row.get("车次") or not row.get("上车站") or not row.get("下车站"):
            continue
        origin = station(row["上车站"], row.get("上车城市"), row.get("上车国家"), foreign_stations)
        destination = station(row["下车站"], row.get("下车城市"), row.get("下车国家"), foreign_stations)
        day_offset = inferred_day_offset(row.get("日期"), row.get("上车时间"), row.get("下车时间"), duration_minutes(row), origin["timezone"], destination["timezone"])
        origin["utcOffset"] = utc_offset(row.get("日期"), row.get("上车时间"), origin["timezone"])
        destination["utcOffset"] = utc_offset(row.get("日期"), row.get("下车时间"), destination["timezone"], day_offset)
        catalog.append({
            "id": f"foreign-rail-{index}", "mode": "train", "date": row.get("日期"),
            "sourceNumber": row.get("序号"),
            "departureDate": row.get("日期"), "arrivalDate": shifted_date(row.get("日期"), day_offset),
            "dayOffset": day_offset,
            "service": str(row["车次"]), "operator": row.get("局属"),
            "origin": origin, "destination": destination,
            "startTerminal": None, "endTerminal": None,
            "departureTime": row.get("上车时间"), "arrivalTime": row.get("下车时间"),
            "seat": None, "cabin": None,
            "durationMinutes": duration_minutes(row), "distanceKm": row.get("里程"), "speedKmh": row.get("速度"),
            "model": row.get("车型"), "rarity": None, "routeId": None,
        })

    for index, row in enumerate(rows("飞机乘坐记录.xlsx", "Sheet1"), 1):
        if not row.get("航班号") or not row.get("start") or not row.get("depart"):
            continue
        code = row.get("pic")
        logo = f"pic/{code}.png" if code and (ROOT / "pic" / f"{code}.png").exists() else None
        origin = airport(row["start"], row.get("起飞"), row.get("起飞EN"), row.get("起飞城市"), row.get("起飞国家"), province_by_city, row.get("日期"), row.get("起飞时间"))
        destination = airport(row["depart"], row.get("降落"), row.get("降落EN"), row.get("降落城市"), row.get("降落国家"), province_by_city, row.get("日期"), row.get("降落时间"))
        day_offset = row.get("跨天") if isinstance(row.get("跨天"), int) else inferred_day_offset(row.get("日期"), row.get("起飞时间"), row.get("降落时间"), duration_minutes(row), origin["timezone"], destination["timezone"])
        destination["utcOffset"] = utc_offset(row.get("日期"), row.get("降落时间"), destination["timezone"], day_offset)
        catalog.append({
            "id": f"air-{index}", "mode": "air", "date": row.get("日期"),
            "sourceNumber": row.get("序号"),
            "departureDate": row.get("日期"), "arrivalDate": shifted_date(row.get("日期"), day_offset),
            "dayOffset": day_offset,
            "service": str(row["航班号"]), "airline": row.get("航空公司") or row.get("airlines"),
            "logo": logo,
            "origin": origin,
            "destination": destination,
            "departureTime": row.get("起飞时间"), "arrivalTime": row.get("降落时间"),
            "durationMinutes": duration_minutes(row), "distanceKm": row.get("飞行距离"),
            "seat": row.get("座位号"), "cabin": row.get("席位"),
            "model": row.get("机型"),
            "rarity": plane_rarity.get(row.get("type")), "registration": row.get("飞机编号"),
        })

    # Undated trips precede dated trips and follow the workbook's 序号 column.
    # Stable sorting keeps source order for rows without a numeric 序号.
    catalog.sort(key=lambda row: (
        row["date"] is not None,
        row["date"] or "",
        row["sourceNumber"] if row["date"] is None and isinstance(row["sourceNumber"], (int, float)) else float("inf"),
    ))
    OUTPUT.write_text(json.dumps(catalog, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Wrote {len(catalog)} trips to {OUTPUT}")


if __name__ == "__main__":
    main()
